import streamlit as st
import datetime
import pandas as pd
import numpy as np
import os
from io import BytesIO

# para ler .docx e .odt
from docx import Document
from odf.opendocument import load as load_odf
from odf import text, teletype

# para gerar PDF
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter

# módulos de dados
from pacientes import adicionar_paciente, obter_pacientes
from agendamentos import adicionar_agendamento, obter_agendamentos
from prontuario import adicionar_prontuario, obter_prontuarios_por_paciente
from financeiro import adicionar_transacao, obter_transacoes
from comunicacao import adicionar_comunicacao
from laudos import adicionar_laudo, obter_laudos
from relatorios import gerar_relatorio, relatorio_por_tipo_agendamento

# UI
from streamlit_option_menu import option_menu
from st_aggrid import AgGrid, GridOptionsBuilder

# --- Configuração da página ---
st.set_page_config(
    page_title="Consultório de Neuropediatria",
    page_icon="logo.png",
    layout="wide"
)

# --- Helpers de template ---
def load_docx(path: str) -> str:
    doc = Document(path)
    return "\n".join(p.text for p in doc.paragraphs)

def load_odt(path: str) -> str:
    odt = load_odf(path)
    paras = odt.getElementsByType(text.P)
    return "\n".join(teletype.extractText(p) for p in paras)

# --- Componentes de interface ---
def mostrar_logo():
    st.image("logo.png", width=150)
    st.markdown("---")

def exibir_tabela(df: pd.DataFrame):
    if df is None or df.empty:
        st.info("Nenhum dado para exibir.")
        return
    df.columns = df.columns.map(str)
    gb = GridOptionsBuilder.from_dataframe(df)
    gb.configure_pagination(paginationAutoPageSize=True)
    gb.configure_default_column(filter=True, sortable=True)
    AgGrid(df, gridOptions=gb.build(), update_mode="MODEL_CHANGED")

# --- Páginas ---
def pagina_inicial():
    st.title("Início")
    st.write("Bem‑vindo ao sistema de gestão do consultório de neuropediatria.")

def cadastro_pacientes():
    st.title("Cadastro de Pacientes")
    with st.form("form_cadastro"):
        nome = st.text_input("Nome")
        data_nasc = st.date_input("Data de Nascimento")
        idade = st.number_input("Idade", min_value=0, max_value=150)
        cpf = st.text_input("CPF")
        rg = st.text_input("RG")
        email = st.text_input("E‑mail")
        tel = st.text_input("Telefone")
        tel2 = st.text_input("Telefone Adicional")
        endereco = st.text_input("Endereço")
        numero = st.text_input("Número")
        comp = st.text_input("Complemento")
        bairro = st.text_input("Bairro")
        cep = st.text_input("CEP")
        cidade = st.text_input("Cidade")
        estado = st.text_input("Estado")
        plano = st.text_input("Plano de Saúde")
        historico = st.text_area("Histórico Médico")
        obs = st.text_area("Observações")
        enviar = st.form_submit_button("Salvar")
    if enviar:
        adicionar_paciente(
            nome,
            data_nasc.isoformat(),
            idade,
            cpf, rg, email,
            tel, tel2,
            endereco, numero, comp, bairro, cep, cidade, estado,
            plano, historico, obs
        )
        st.success("Paciente cadastrado com sucesso!")

def agendamentos_pagina():
    st.title("Agendamento de Consultas")
    pacientes = obter_pacientes()
    if not pacientes:
        st.info("Cadastre um paciente primeiro!")
        return
    op = {f"{p['id']} ‑ {p['nome']}": p['id'] for p in pacientes}
    sel = st.selectbox("Paciente", list(op.keys()))
    pid = op[sel]
    with st.form("form_agendamento"):
        data_c = st.date_input("Data da Consulta", datetime.date.today())
        hora_c = st.time_input("Hora da Consulta", datetime.time(9,0))
        obs = st.text_area("Observações")
        tipo = st.selectbox("Tipo de Consulta", ["Plano de Saúde","Particular"])
        enviar = st.form_submit_button("Agendar")
    if enviar:
        adicionar_agendamento(pid, data_c.isoformat(), hora_c.strftime("%H:%M"), obs, tipo)
        st.success("Consulta agendada com sucesso!")

def prontuario_pagina():
    st.title("Prontuário Eletrônico")
    pacientes = obter_pacientes()
    if not pacientes:
        st.info("Cadastre um paciente primeiro!")
        return
    op = {f"{p['id']} ‑ {p['nome']}": p['id'] for p in pacientes}
    sel = st.selectbox("Paciente", list(op.keys()))
    pid = op[sel]
    with st.form("form_prontuario"):
        desc = st.text_area("Descrição de Atendimento")
        enviar = st.form_submit_button("Salvar")
    if enviar:
        adicionar_prontuario(pid, desc, datetime.date.today().isoformat())
        st.success("Registro salvo com sucesso!")
    regs = obter_prontuarios_por_paciente(pid)
    exibir_tabela(pd.DataFrame(regs))

def financeiro_pagina():
    st.title("Gestão Financeira")
    pacientes = obter_pacientes()
    if not pacientes:
        st.info("Cadastre um paciente primeiro!")
        return
    op = {f"{p['id']} ‑ {p['nome']}": p['id'] for p in pacientes}
    sel = st.selectbox("Paciente", list(op.keys()))
    pid = op[sel]
    with st.form("form_financeiro"):
        data_f = st.date_input("Data da Transação", datetime.date.today())
        valor = st.number_input("Valor", min_value=0.0, step=0.1)
        desc = st.text_area("Descrição")
        enviar = st.form_submit_button("Registrar")
    if enviar:
        adicionar_transacao(pid, data_f.isoformat(), valor, desc)
        st.success("Transação registrada com sucesso!")
    txs = obter_transacoes()
    exibir_tabela(pd.DataFrame(txs))

def comunicacao_pagina():
    st.title("Comunicação com Pacientes")
    pacientes = obter_pacientes()
    if not pacientes:
        st.info("Cadastre um paciente primeiro!")
        return
    op = {f"{p['id']} ‑ {p['nome']}": p['id'] for p in pacientes}
    sel = st.selectbox("Paciente", list(op.keys()))
    pid = op[sel]
    with st.form("form_comunicacao"):
        msg = st.text_area("Mensagem")
        enviar = st.form_submit_button("Enviar")
    if enviar:
        adicionar_comunicacao(pid, msg, datetime.datetime.now().isoformat())
        st.success("Mensagem enviada com sucesso!")

def relatorios_pagina():
    st.title("Relatórios e Estatísticas")
    tot_p, tot_a, tot_f = gerar_relatorio()
    st.write(f"**Pacientes:** {tot_p}")
    st.write(f"**Consultas:** {tot_a}")
    st.write(f"**Financeiro:** R$ {tot_f:.2f}")
    st.subheader("Consultas por Tipo")
    for t, q in relatorio_por_tipo_agendamento().items():
        st.write(f"{t}: {q} consulta(s)")

def dashboard_pagina():
    st.title("Dashboard Completo")
    tot_p, tot_a, tot_f = gerar_relatorio()
    pacs = obter_pacientes()
    media = np.mean([p['idade'] for p in pacs if p.get('idade') is not None]) if pacs else 0
    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Pacientes", tot_p)
    c2.metric("Consultas", tot_a)
    c3.metric("Financeiro", f"R$ {tot_f:.2f}")
    c4.metric("Média de Idade", f"{media:.1f}")
    st.markdown("---")
    st.subheader("Catálogo de Pacientes")
    exibir_tabela(pd.DataFrame(pacs))

def buscar_paciente_pagina():
    st.title("Buscar Paciente")
    nome = st.text_input("Digite parte do nome")
    if nome:
        pacs = obter_pacientes()
        matches = [p for p in pacs if nome.lower() in p["nome"].lower()]
        if matches:
            df = pd.DataFrame(matches)
            exibir_tabela(df)
            op = {f"{p['id']} ‑ {p['nome']}": p['id'] for p in matches}
            sel = st.selectbox("Ver histórico de", list(op.keys()))
            pid = op[sel]
            st.subheader("Histórico de Atendimentos")
            regs = obter_prontuarios_por_paciente(pid)
            exibir_tabela(pd.DataFrame(regs))
        else:
            st.warning("Nenhum paciente encontrado.")

def laudos_pagina():
    st.title("Emissão de Laudo em PDF")
    pacientes = obter_pacientes()
    if not pacientes:
        st.info("Cadastre um paciente primeiro!")
        return
    op = {f"{p['id']} ‑ {p['nome']}": p for p in pacientes}
    sel = st.selectbox("Paciente", list(op.keys()))
    pac = op[sel]
    md = "templates"
    arquivos = [f for f in os.listdir(md) if f.lower().endswith((".docx", ".odt"))]
    if not arquivos:
        st.error("Coloque seus modelos em 'templates/'")
        return
    modelo = st.selectbox("Modelo de Laudo", arquivos)
    path = os.path.join(md, modelo)
    texto = load_docx(path) if modelo.lower().endswith(".docx") else load_odt(path)
    st.markdown("**Edite o texto. Use `{nome}` e `{data}` como placeholders.**")
    laudo_txt = st.text_area("Conteúdo do Laudo", value=texto, height=300)
    if st.button("Gerar PDF"):
        ctx = {"nome": pac["nome"], "data": datetime.date.today().strftime("%d/%m/%Y")}
        rend = laudo_txt.format(**ctx)
        buf = BytesIO()
        c = canvas.Canvas(buf, pagesize=letter)
        t = c.beginText(40, 750)
        for line in rend.split("\n"):
            t.textLine(line)
        c.drawText(t); c.showPage(); c.save()
        buf.seek(0)
        fn = f"laudo_{pac['nome'].replace(' ','_')}_{ctx['data'].replace('/','-')}.pdf"
        st.download_button("Download PDF", buf, file_name=fn, mime="application/pdf")

# --- Main ---
def main():
    mostrar_logo()
    selected = option_menu(
        menu_title=None,
        options=[
            "Início","Dashboard","Buscar Paciente","Cadastro de Pacientes",
            "Agendamentos","Prontuário","Financeiro","Comunicação",
            "Relatórios","Laudos"
        ],
        icons=[
            "house","bar-chart","search","person-plus",
            "calendar","file-text","wallet","chat-dots",
            "clipboard-data","file-earmark-text"
        ],
        orientation="horizontal",
        default_index=0
    )
    if selected == "Início":
        pagina_inicial()
    elif selected == "Dashboard":
        dashboard_pagina()
    elif selected == "Buscar Paciente":
        buscar_paciente_pagina()
    elif selected == "Cadastro de Pacientes":
        cadastro_pacientes()
    elif selected == "Agendamentos":
        agendamentos_pagina()
    elif selected == "Prontuário":
        prontuario_pagina()
    elif selected == "Financeiro":
        financeiro_pagina()
    elif selected == "Comunicação":
        comunicacao_pagina()
    elif selected == "Relatórios":
        relatorios_pagina()
    elif selected == "Laudos":
        laudos_pagina()

if __name__ == "__main__":
    main()
